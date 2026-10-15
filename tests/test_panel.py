"""Admin panel (handlers/panel.py): /find, ban, qaytarildi deb yopish, suyunchilar, /export, /broadcast."""
import asyncio
from io import BytesIO

from aiogram.methods import CopyMessage, SendDocument
from openpyxl import load_workbook

import database as db
from config import ADMIN_GROUP_ID
from tests.fakes import ADMIN, buttons, make_claim, make_item

FINDER, OWNER = 20, 30


async def setup_item():
    await db.save_user(FINDER, "Topuvchi", "topuvchi")
    await db.save_user(OWNER, "Egasi", "egasi")
    item = await make_item()
    claim = await make_claim(item["id"])
    return item, claim


async def test_help_only_in_admin_group(tg):
    await tg.send("/help")
    assert "/find" in tg.last_text() and "/broadcast" in tg.last_text()
    await tg.send("/stats", user_id=FINDER)
    assert not any("Statistika" in text for text in tg.texts(FINDER))


async def test_find_item_and_claim(tg):
    item, claim = await setup_item()
    for query in ("T0001", "1", "#T0001", f"m{claim['id']}", f"№M{claim['id']}"):
        await tg.send(f"/find {query}")
        card = tg.last_text()
        assert "#T0001" in card and "Kanalda" in card and f"№{claim['id']}" in card, query
    assert "8600 1111" in card  # admin karta raqamini ko'radi
    assert f"fi:ret:{item['id']}" in buttons(tg.sent()[-1])


async def test_find_user_and_phone(tg):
    await setup_item()
    await tg.send("/find @topuvchi")
    assert "Topuvchi" in tg.last_text() and "#T0001" in tg.last_text()
    await tg.send("/find 55555555")
    assert "hech narsa topilmadi" in tg.last_text()
    await tg.send("/find +998 93 765-43-21")
    assert "Egasi sifatida" in tg.last_text()
    await tg.send("/find 901234567")
    assert "Topuvchi sifatida" in tg.last_text()


async def test_ban_blocks_everything_and_unban_restores(tg):
    await setup_item()
    await tg.send("/ban @egasi firibgarlik")
    assert "bloklandi" in tg.last_text() and await db.is_banned(OWNER)
    assert any("firibgarlik" in text for text in tg.texts(OWNER))

    await tg.send("/start", user_id=OWNER)
    assert tg.texts(OWNER) == ["⛔ Siz botdan foydalanishdan cheklangansiz."]
    await tg.press("claim:1", user_id=OWNER)
    assert tg.alerts() and not tg.sent()

    await tg.send(f"/ban {ADMIN}")
    assert "O'zingizni" in tg.last_text()
    await tg.press(f"ban:off:{OWNER}")
    await tg.send("/start", user_id=OWNER)
    assert "Assalomu alaykum" in tg.texts(OWNER)[0]


async def test_mark_returned_needs_confirmation(tg):
    item, claim = await setup_item()
    await tg.press(f"fi:ret:{item['id']}")
    assert (await db.get_item(item["id"]))["status"] == "published"  # birinchi bosish — faqat tasdiq so'raydi
    await tg.press(f"fi:ret2:{item['id']}")
    assert (await db.get_item(item["id"]))["status"] == "returned"
    assert (await db.get_claim(claim["id"]))["status"] == "rejected"
    assert tg.texts(FINDER) and tg.texts(OWNER)


async def test_payouts(tg):
    item = await make_item(status="returned")
    claim = await make_claim(item["id"])
    await db.update_claim(claim["id"], status="completed", handover="confirmed")
    await tg.send("/payouts")
    assert "1 ta, jami 25 000 so'm" in tg.last_text() and "8600 1111" in tg.last_text()
    await tg.press(f"po:list:{claim['id']}")
    assert any("suyunchingiz" in text for text in tg.texts(FINDER))
    await tg.press(f"po:one:{claim['id']}")
    assert tg.alerts() and not tg.sent()  # ikkinchi marta belgilab bo'lmaydi
    await tg.send("/payouts")
    assert "Barcha suyunchilar o'tkazilgan" in tg.last_text()


async def test_export_excel(tg):
    await setup_item()
    await tg.send("/export")
    workbook = load_workbook(BytesIO(tg.to(ADMIN_GROUP_ID, SendDocument)[0].document.data))
    assert workbook.sheetnames == ["E'lonlar", "Murojaatlar", "Qidiruvlar", "Foydalanuvchilar"]
    assert workbook["E'lonlar"]["A2"].value == "#T0001"
    assert workbook["Foydalanuvchilar"].max_row == 3


async def test_broadcast(tg):
    await setup_item()
    await db.save_user(999, "Blokladi", None)
    tg.blocked.add(999)  # botni bloklagan
    await tg.send("/broadcast")
    assert "javob (reply)" in tg.last_text()

    target = tg.message("Yangilik!", message_id=77)
    await tg.send("/broadcast", reply_to=target)
    assert "3</b> ta foydalanuvchiga" in tg.last_text()
    await db.set_banned(OWNER, True)
    await tg.press("bc:go:77")
    await asyncio.sleep(0.5)
    assert sorted(call.chat_id for call in tg.sent(CopyMessage)) == [FINDER, 999]  # bloklangan egasiga emas
    assert any("Yetkazildi: 1" in text and "Yetmadi: 1" in text for text in tg.edits())
