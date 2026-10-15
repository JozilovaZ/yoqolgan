"""Bazani Excel (.xlsx) faylga chiqarish — admin /export buyrug'i uchun."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

import database as db
from utils import CLAIM_STATUSES, HANDOVER_STATUSES, ITEM_STATUSES, category_label, item_code, local_time, lost_code

HEADER_FILL = PatternFill("solid", fgColor="DDE7F0")


def _sheet(workbook: Workbook, title: str, headers: list[str], rows: list[list]) -> None:
    sheet = workbook.create_sheet(title)
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, header in enumerate(headers, start=1):
        longest = max([len(str(row[index - 1] or "")) for row in rows] + [len(header)])
        sheet.column_dimensions[get_column_letter(index)].width = min(max(longest + 2, 8), 50)


async def build_export() -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)

    items = await db.export_table("items")
    _sheet(
        workbook,
        "E'lonlar",
        ["Raqam", "Holati", "Turi", "Viloyat", "Joy", "Topilgan vaqt", "Izoh",
         "Topuvchi", "Username", "Telefon", "Karta", "Berilgan", "Yopilgan"],
        [
            [
                f"#{item_code(i['id'])}", ITEM_STATUSES.get(i["status"], i["status"]), category_label(i["category"]),
                i["region"], i["place"], i["found_date"], i["description"],
                i["finder_name"], i["finder_username"], i["finder_phone"], i["finder_card"],
                local_time(i["created_at"]), local_time(i["closed_at"]),
            ]
            for i in items
        ],
    )

    claims = await db.export_table("claims")
    _sheet(
        workbook,
        "Murojaatlar",
        ["Murojaat", "E'lon", "Holati", "Topshirish", "Egasi", "Username", "Telefon",
         "To'lov ID", "Suyunchi o'tkazildi", "Berilgan"],
        [
            [
                c["id"], f"#{item_code(c['item_id'])}", CLAIM_STATUSES.get(c["status"], c["status"]),
                HANDOVER_STATUSES.get(c["handover"], "") if c["status"] == "completed" else "",
                c["owner_name"], c["owner_username"], c["owner_phone"],
                c["payment_charge_id"], local_time(c["payout_at"]) if c["payout_at"] else "",
                local_time(c["created_at"]),
            ]
            for c in claims
        ],
    )

    reports = await db.export_table("lost_reports")
    _sheet(
        workbook,
        "Qidiruvlar",
        ["Raqam", "Holati", "Turi", "Viloyat", "Joy", "Yo'qotilgan kun", "Tavsif", "Egasi", "Username", "Berilgan"],
        [
            [
                f"#{lost_code(r['id'])}", "🔔 Faol" if r["status"] == "active" else "Yopilgan",
                category_label(r["category"]), r["region"], r["place"], r["lost_day"], r["description"],
                r["owner_name"], r["owner_username"], local_time(r["created_at"]),
            ]
            for r in reports
        ],
    )

    users = await db.export_table("users")
    _sheet(
        workbook,
        "Foydalanuvchilar",
        ["ID", "Ism", "Username", "Bloklangan", "Ofertaga rozi (tahrir)", "Rozilik vaqti", "Qo'shilgan"],
        [
            [
                u["user_id"], u["full_name"], u["username"], "⛔ Ha" if u["banned"] else "",
                u["terms_version"] or "", local_time(u["terms_accepted_at"]) if u["terms_accepted_at"] else "",
                local_time(u["created_at"]),
            ]
            for u in users
        ],
    )

    result = BytesIO()
    workbook.save(result)
    return result.getvalue()
